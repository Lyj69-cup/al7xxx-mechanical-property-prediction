# -*- coding: utf-8 -*-
"""Supplementary Table S1 (v4): source summary + per-record provenance."""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
DATA = str(ROOT/"data"); OUT = str(ROOT/"supplementary")
os.makedirs(OUT, exist_ok=True)

df = pd.read_csv(os.path.join(DATA, "al7xxx_dataset_v4.csv"), encoding="utf-8-sig")

# ---- source-level metadata (citations verified against the published items) ----
META = {
 "BHAT2023": dict(
   cite="Bhat, N.; Barnard, A.; Birbilis, N. Aluminium alloy dataset for supervised learning. Mendeley Data, V1, 2023.",
   doi="10.17632/b6br4yk6r3.1", url="https://data.mendeley.com/datasets/b6br4yk6r3/1",
   lic="CC BY 4.0",
   prov="Open curated dataset. The 7xxx subset was selected by an objective composition rule (Zn >= 3 wt% and Zn > Cu), not by the dataset's 'class' label. Each record is traceable to a source row index; processing-class text is retained verbatim. No per-record temperatures/times are reported in this source."),
 "XU2015": dict(
   cite="Tan, C.; Xu, X.-J.; Jiang, W.; Sun, L.-S.; Zhang, X.-L.; Fan, Y.-Z.; Zhao, J.-J. Effects of pre-recovery and multi-stage solution treatment on an Al-Zn-Mg-Cu alloy. The Chinese Journal of Nonferrous Metals, 2015, 25(11), 3019-3025. (in Chinese)",
   doi="-", url="http://www.ysxbcn.com/down/2015/11_cn/08-p3019-33029.pdf",
   lic="Subscription journal; measured points used for scholarly verification",
   prov="Manually tabulated from the published tensile-property tables. Includes pre-recovery, multi-stage solution and T652 (2% pre-compression) routes. Yield strength not reported for these two records."),
 "XU2016": dict(
   cite="Zhang, X.-L.; Xu, X.-J.; Ling, Z.-Y.; Jiang, W.; Tan, C.; Zhao, J.-J. Effect of pre-deformation on microstructure and mechanical properties of an Al-Zn-Mg-Cu alloy. The Chinese Journal of Nonferrous Metals, 2016, 26(3), 507-513. (in Chinese)",
   doi="-", url="http://www.ysxbcn.com/down/2016/03_cn/04-p0507-33268.pdf",
   lic="Subscription journal; measured points used for scholarly verification",
   prov="Manually tabulated from the published tensile-property tables; peak-aged T6 with and without 2-3% pre-stretch."),
 "XU2018": dict(
   cite="Zhang, C.; Xu, X.-J.; Zhang, J.; Huang, P.; Du, D.-H.; Jia, W.-J. Effect of heating rate during solution treatment on an Al-Zn-Mg-Cu alloy. Materials for Mechanical Engineering, 2018, 42(8), 24-27. (in Chinese)",
   doi="10.11973/jxgccl201808005", url="https://jxgccl.mat-test.com/cn/article/pdf/preview/10.11973/jxgccl201808005.pdf",
   lic="Subscription journal; measured points used for scholarly verification",
   prov="Manually tabulated from the published tensile-property tables; slow (3.6 C/h) vs fast (180 C/h) heating to the solution temperature."),
 "LI2020": dict(
   cite="Li, J.; Zhang, Y.; Cao, X.; Zeng, Q.; Zhuang, Y.; Qian, X.; Chen, H. Accelerated discovery of high-strength aluminum alloys by machine learning. Communications Materials, 2020, 1, 73.",
   doi="10.1038/s43246-020-00074-2", url="https://www.nature.com/articles/s43246-020-00074-2",
   lic="CC BY 4.0 (open access)",
   prov="Tabulated from Table 1: twelve iteratively designed new alloys in the as-cast, screened condition. Only tensile strength is reported. The 'EI value' column in the source is the expected-improvement acquisition function, NOT elongation; it was not used as a target."),
 "PMC9571267": dict(
   cite="Zhang, W. et al. Effects of forging temperature on microstructure and mechanical properties of an Al-Zn-Mg-Cu alloy, 2022.",
   doi="-", url="https://pmc.ncbi.nlm.nih.gov/articles/PMC9571267/",
   lic="Open access (PMC); reuse terms per the article",
   prov="Tabulated from the published tensile results: as-forged and three forging temperatures (450, 250, -180 C), followed by solution and two-step aging. Full author list and journal metadata are given in the manuscript reference list."),
}
ORDER = ["BHAT2023", "XU2015", "XU2016", "XU2018", "LI2020", "PMC9571267"]

# ---- styles ----
HEAD = PatternFill("solid", fgColor="17365D")
SUBHEAD = PatternFill("solid", fgColor="2E5496")
HF = Font(name="Arial", bold=True, color="FFFFFF", size=10.5)
TF = Font(name="Arial", size=9.5)
thin = Side(style="thin", color="B8C7D9")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
LAL = Alignment(horizontal="left", vertical="center", wrap_text=True)
CAC = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_header(ws, row, ncol, fill=HEAD):
    for c in range(1, ncol+1):
        cell = ws.cell(row, c); cell.fill = fill; cell.font = HF
        cell.alignment = CAC; cell.border = BORD

wb = Workbook()

# ===== Sheet 1: source summary =====
ws = wb.active; ws.title = "Source summary"
cols = ["Source ID", "Records", "n (yield)", "n (tensile)", "n (elongation)",
        "Full citation", "DOI", "URL", "License", "Extraction and provenance"]
ws.append(cols); style_header(ws, 1, len(cols))
for sid in ORDER:
    d = df[df.source_id == sid]
    ws.append([sid, len(d), int(d.YS_MPa.notna().sum()), int(d.UTS_MPa.notna().sum()),
               int(d.EL_pct.notna().sum()), META[sid]["cite"], META[sid]["doi"],
               META[sid]["url"], META[sid]["lic"], META[sid]["prov"]])
# totals
tot = ["TOTAL", len(df), int(df.YS_MPa.notna().sum()), int(df.UTS_MPa.notna().sum()),
       int(df.EL_pct.notna().sum()), "6 sources (1 open dataset + 5 peer-reviewed studies)",
       "", "", "", "Every record is traceable to a DOI or URL; no synthetic or interpolated records."]
ws.append(tot)
widths = [13, 9, 9, 10, 12, 46, 24, 40, 30, 60]
for i, wd in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = wd
for r in range(2, ws.max_row+1):
    for c in range(1, len(cols)+1):
        cell = ws.cell(r, c); cell.font = TF; cell.border = BORD
        cell.alignment = CAC if c in (1,2,3,4,5) else LAL
last = ws.max_row
for c in range(1, len(cols)+1):
    ws.cell(last, c).font = Font(name="Arial", bold=True, size=9.5)
    ws.cell(last, c).fill = PatternFill("solid", fgColor="DCE6F1")
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{last-1}"

# ===== Sheet 2: per-record provenance =====
ws2 = wb.create_sheet("Record provenance")
c2 = ["Record ID", "Source", "Temper", "Product form", "Data origin",
      "Localization (table / row)", "Zn (wt%)", "Mg (wt%)", "Cu (wt%)",
      "Yield (MPa)", "Tensile (MPa)", "Elong. (%)", "DOI / URL"]
ws2.append(c2); style_header(ws2, 1, len(c2), fill=SUBHEAD)
for _, r in df.iterrows():
    ws2.append([r.record_id, r.source_id, r.temper, r.product_form, r.data_origin,
                r.localization, r.Zn, r.Mg, r.Cu,
                None if pd.isna(r.YS_MPa) else r.YS_MPa,
                None if pd.isna(r.UTS_MPa) else r.UTS_MPa,
                None if pd.isna(r.EL_pct) else r.EL_pct,
                (r.doi if str(r.doi) not in ("-","nan") else r.url)])
w2 = [11, 12, 20, 13, 26, 30, 9, 9, 9, 11, 11, 10, 42]
for i, wd in enumerate(w2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = wd
numcols = (1,2,3,4,5,6)
for r in range(2, ws2.max_row+1):
    for c in range(1, len(c2)+1):
        cell = ws2.cell(r, c); cell.font = TF; cell.border = BORD
        cell.alignment = CAC if c in (1,2,7,8,9,10,11,12) else LAL
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = f"A1:M{ws2.max_row}"

path = os.path.join(OUT, "Supplementary_Table_S1_Data_Sources_v4.xlsx")
wb.save(path)
print("saved", path)
print("sources:", df.source_id.value_counts().to_dict())
print("totals n:", len(df), int(df.YS_MPa.notna().sum()),
      int(df.UTS_MPa.notna().sum()), int(df.EL_pct.notna().sum()))
